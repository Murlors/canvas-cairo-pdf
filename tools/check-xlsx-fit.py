"""源单元格完整性回归：不同默认字体、列宽范围和最后一列宽度。"""
import json,subprocess,tempfile
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
root=Path(__file__).resolve().parent.parent
out=Path(tempfile.mkdtemp(prefix='xlsx-fit-regression-',dir=root/'output'))
results=[]
cases=[('Arial',10,12),('Times New Roman',14,24),('Courier New',11,40),('MissingFontExample',12,32)]
for family,size,columns in cases:
    book=Workbook();book._fonts[0]=Font(name=family,sz=size);sheet=book.active
    sheet.column_dimensions.group('B',get_column_letter(columns-1));sheet.column_dimensions['B'].width=18
    sheet.column_dimensions[get_column_letter(columns)].width=35
    expected=[]
    for r in range(1,13):
        for c in range(1,columns+1):
            value=f'R{r:02}C{c:02}';expected.append(value);sheet.cell(r,c,value)
    source=out/f'{columns}.xlsx';book.save(source)
    process=subprocess.run(['node','scripts/run-pages.ts',str(source)],cwd=root,capture_output=True,text=True,check=True)
    info=json.loads(process.stdout);directory=Path(info['run'])/'pages';pdf=PdfReader(directory/'document.pdf')
    text=''.join(page.extract_text() for page in pdf.pages)
    missing=[value for value in expected if text.count(value)!=1]
    assert not missing,(family,missing)
    row={'font':family,'size':size,'columns':columns,'cells':len(expected),'missingOrDuplicate':0,'pages':len(pdf.pages),'run':str(directory)}
    results.append(row);print(json.dumps(row),flush=True)
    (out/'results.json').write_text(json.dumps(results,indent=2))
print(out)
