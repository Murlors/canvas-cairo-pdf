"""补充常见办公结构；合成回归样本，不作为 Word/WPS 排版金标准。"""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path(__file__).resolve().parent.parent / 'fixtures/representative'
out.mkdir(parents=True, exist_ok=True)
doc = Document()
style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(11)
style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'PingFang SC')
doc.sections[0].header.paragraphs[0].text = 'Quarterly Operations / 季度运营报告'
footer = doc.sections[0].footer.paragraphs[0]
footer.add_run('Page ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
doc.add_heading('季度项目执行报告', 0)
doc.add_paragraph('SUMMARY-001 本文用于检查常见办公文档，不含真实个人数据。')
for i in range(1, 4):
    doc.add_heading(f'{i}. 项目进展 / Project progress', 1)
    for j in range(1, 6):
        doc.add_paragraph(f'ITEM-{i}-{j} 预算 ¥12,500.00，完成率 85%。中文与 English 混排；检查换行、标点和项目列表。', style='List Bullet')
doc.add_heading('跨页明细表', 1)
table = doc.add_table(rows=1, cols=4); table.style = 'Table Grid'
for cell, text in zip(table.rows[0].cells, ['编号', '项目', '负责人', '金额']): cell.text = text
repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
for i in range(1, 46):
    for cell, text in zip(table.add_row().cells, [f'ROW-{i:03}', '本地文档处理 / Local processing', 'Team A', f'{i * 120}.00']): cell.text = text
row = table.add_row(); row.cells[0].merge(row.cells[2]).text = '合并单元格 / Total'; row.cells[3].text = '124200.00'
doc.add_paragraph('REPORT-END-001 文档结束。')
doc.save(out / 'office-report.docx')
print(out / 'office-report.docx')
