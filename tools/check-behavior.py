"""分页/选项验收及真实worker路径对照；不比较字重细节。"""
import json,os,subprocess,tempfile
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from PIL import Image,ImageChops,ImageStat
root=Path(__file__).resolve().parent.parent
fixtures=root/'fixtures/representative'
fixtures.mkdir(parents=True, exist_ok=True)
w=Workbook();s=w.active;s.title='宽表'
for c in range(1,33):
    s.column_dimensions[get_column_letter(c)].width=20
    for r in range(1,81):s.cell(r,c,f'R{r:03}C{c:02}')
h=w.create_sheet('隐藏');h['A1']='HIDDEN-ONLY';h.sheet_state='hidden'
w.save(fixtures/'wide.xlsx')
(fixtures/'long.md').write_text('# 多页笔记\n\n'+'\n\n'.join(f'MD-{i:03} 中文内容与 English text，逐段核对分页和末尾内容。' for i in range(1,91)))
out=Path(tempfile.mkdtemp(prefix='behavior-check-',dir=root/'output'));rows=[]
def run(source,**options):
    env={**os.environ,**options}
    p=subprocess.run(['node','scripts/run-pages.ts',str(fixtures/source)],cwd=root,env=env,capture_output=True,text=True)
    if p.returncode:raise RuntimeError(p.stderr)
    info=json.loads(p.stdout);directory=Path(info['run'])/'pages'
    return directory,PdfReader(directory/'document.pdf')
for name,source,opts in [('wide-fit','wide.xlsx',{}),('wide-actual','wide.xlsx',{'PLIFLO_XLSX_SCALE':'actual'}),('hidden-selected','wide.xlsx',{'PLIFLO_XLSX_SHEET':'1'}),('markdown','long.md',{})]:
    directory,reader=run(source,**opts);text=''.join(p.extract_text() for p in reader.pages)
    if name=='hidden-selected':assert 'HIDDEN-ONLY' in text
    elif name.startswith('wide'):assert 'HIDDEN-ONLY' not in text
    else:assert all(text.count(f'MD-{i:03}')==1 for i in range(1,91))
    row={'case':name,'pages':len(reader.pages),'run':str(directory),'characters':len(text)};rows.append(row)
    if name=='wide-fit':
        row['missingOrDuplicateCells']=[f'R{r:03}C{c:02}' for r in range(1,81) for c in range(1,33) if text.count(f'R{r:03}C{c:02}')!=1]
        row['contentComplete']=not row['missingOrDuplicateCells']
    if name in ['wide-fit','wide-actual']:
        reference,worker=run(source,PLIFLO_WORKER_REFERENCE='1',**opts)
        assert len(worker.pages)==len(reader.pages)
        diffs=[]
        for i,(a,b) in enumerate(zip(reader.pages,worker.pages),1):
            assert a.mediabox==b.mediabox
            left=Image.open(directory/f'page-{i:03}.png').convert('RGB');right=Image.open(reference/f'page-{i:03}.png').convert('RGB')
            assert left.size==right.size
            diffs.append(sum(ImageStat.Stat(ImageChops.difference(left,right)).mean)/3)
        row['workerReference']=str(reference);row['mainWorkerCanvasMAE']=diffs
    (out/'results.json').write_text(json.dumps(rows,indent=2));print(json.dumps(row),flush=True)
assert rows[0]['pages']<rows[1]['pages'],'Wide fit must scale and reduce page count'
print(out)
if any(row.get('contentComplete') is False for row in rows):raise SystemExit('FAILED: fit width omitted source cells; see results.json')
