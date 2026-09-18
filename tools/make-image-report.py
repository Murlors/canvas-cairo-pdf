"""可复现的图片负载：程序生成图表，不依赖远程素材或真实用户数据。"""
from pathlib import Path
from io import BytesIO
import random
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Mm

out=Path(__file__).resolve().parent.parent/'fixtures/representative'
out.mkdir(parents=True, exist_ok=True)
doc=Document()
rng=random.Random(20260918)
for page in range(6):
    if page: doc.add_page_break()
    doc.add_heading(f'IMAGE-PAGE-{page+1:02} Project dashboard', 1)
    doc.add_paragraph('图文报告：检查图片、标题和正文在 PDF 中是否完整。')
    for chart in range(2):
        im=Image.new('RGB',(1600,800),'#f1f4f8');draw=ImageDraw.Draw(im)
        for i in range(24):
            h=rng.randint(50,640)
            draw.rectangle((60+i*60,730-h,100+i*60,730),fill=('#387ea3' if i%2 else '#58a18d'))
        draw.line((40,735,1540,735),fill='#222222',width=3)
        buf=BytesIO();im.save(buf,format='PNG');buf.seek(0)
        doc.add_picture(buf,width=Mm(150))
        doc.add_paragraph(f'FIGURE-{page*2+chart+1:02} Independent chart / 图表说明')
doc.save(out/'image-report.docx')
print(out/'image-report.docx')
