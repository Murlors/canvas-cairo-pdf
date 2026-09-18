from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
from openpyxl.styles import PatternFill,Border,Side

out=Path(__file__).resolve().parent.parent/'fixtures/representative'
out.mkdir(parents=True, exist_ok=True)
p=Presentation()
p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
for i in range(3):
    slide=p.slides.add_slide(p.slide_layouts[5]);slide.shapes.title.text=f'SLIDE-{i+1:02} 项目报告 / Project report'
    box=slide.shapes.add_textbox(Inches(1),Inches(1.5),Inches(10),Inches(1));box.text='中文搜索 Mixed English 123 — 每页内容完整'
    table=slide.shapes.add_table(4,3,Inches(1),Inches(3),Inches(10),Inches(2)).table
    for r in range(4):
        for c in range(3):table.cell(r,c).text=f'{r+1}-{c+1} 测试'
p.save(out/'slides.pptx')
w=Workbook();s=w.active;s.title='预算'
for row in range(1,101):
    for col in range(1,15):
        cell=s.cell(row,col,f'R{row:03}C{col:02}' if col>1 else f'项目{row:03}')
        cell.border=Border(bottom=Side(style='thin',color='CCCCCC'))
        if row==1:cell.fill=PatternFill('solid',fgColor='DCEAF2')
s.merge_cells('A102:D102');s['A102']='TOTAL-END 汇总'
w.create_sheet('空表');hidden=w.create_sheet('隐藏');hidden['A1']='HIDDEN-SENTINEL';hidden.sheet_state='hidden'
w.save(out/'workbook.xlsx')
print(out)
