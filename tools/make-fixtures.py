"""可复现的真实 DOCX 验收文件：正文、表格、图片、自动分页及不同 section 尺寸。"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw
import json

ROOT = Path(__file__).resolve().parent.parent / 'fixtures'
ROOT.mkdir(exist_ok=True)
im = Image.new('RGB', (900, 220), '#edf3f7')
d = ImageDraw.Draw(im)
for i, h in enumerate([70, 125, 95, 170, 145, 185]):
    d.rectangle((65+i*130, 205-h, 140+i*130, 205), fill=['#427b98','#58a18d'][i%2])
d.line((35,205,870,205),fill='#344454',width=3)
im.save(ROOT/'chart.png')

def run(p, text, font='Arial', east='PingFang SC', size=11, bold=False):
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), east)
    return r

def section(sec, w=210, h=297):
    sec.page_width=Mm(w); sec.page_height=Mm(h)
    sec.top_margin=Mm(20); sec.bottom_margin=Mm(20)
    sec.left_margin=Mm(22); sec.right_margin=Mm(22)

doc=Document(); section(doc.sections[0])
normal=doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(8)
run(doc.add_paragraph(), 'Pliflo 文档处理评估报告', size=24, bold=True)
run(doc.add_paragraph(), '2026 年 9 月 · 本地渲染实验 / Local rendering assessment', size=12)
run(doc.add_paragraph(), '一、评估摘要', size=16, bold=True)
run(doc.add_paragraph(), '本报告验证中文、中英混排、字体回退、表格与图片在直接 PDF 输出中的保真度。文件仅用于离线渲染测试，不提交任何打印任务。')
run(doc.add_paragraph(), 'SEARCH-CJK-001 中文搜索复制验证：你好世界，批量打印，字体映射。')
run(doc.add_paragraph(), 'Mixed typography: Office 2026 / PDF 1.7 / 价格 ¥128.50 / 完成率 98.5%。')
for font,east,text in [('Arial','PingFang SC','Arial + 苹方：采购清单与预算'),('Times New Roman','Songti SC','Times New Roman + 宋体：季度报告'),('Courier New','Heiti SC','Courier New + 黑体：编号 ABC-0123'),('PlifloMissingFontXYZ','PlifloMissingCJKXYZ','Missing family fallback：缺失字体仍应可见')]:
    run(doc.add_paragraph(), text, font, east)
run(doc.add_paragraph(), '二、实施计划与成本', size=16, bold=True)
table=doc.add_table(rows=1, cols=4); table.style='Table Grid'
for cell,text in zip(table.rows[0].cells,['阶段 Phase','负责人 Owner','预算 Budget','状态 Status']):
    run(cell.paragraphs[0],text,bold=True)
    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCEAF2');cell._tc.get_or_add_tcPr().append(shade)
for row in [('渲染内核','研发 / Engineering','¥12,800','进行中'),('字体诊断','测试 / QA','¥3,600','已安排'),('桌面集成','客户端团队','¥8,200','待验证')]:
    for cell,text in zip(table.add_row().cells,row):run(cell.paragraphs[0],text)
doc.add_picture(str(ROOT/'chart.png'),width=Mm(150))
run(doc.add_paragraph(),'图 1：各阶段模拟工作量 / Illustrative workload',size=9)
doc.add_page_break()
run(doc.add_paragraph(), '三、连续正文与自动分页',size=18,bold=True)
for i in range(1,31):
    run(doc.add_paragraph(),f'段落 {i:02d} / RECORD-{i:03d}  ',bold=True)
    run(doc.paragraphs[-1], '文档分页应由同一个布局引擎确定，每页的边距、行距和文字位置应保持一致。Direct PDF preserves text and vectors while raster output embeds page images. 测试覆盖不同语言与常用标点，确认跨页后没有丢失、重复或顺序颠倒。')
sec=doc.add_section(WD_SECTION_START.NEW_PAGE);section(sec,297,210)
run(doc.add_paragraph(), '四、横向附表 / Landscape appendix',size=20,bold=True)
table=doc.add_table(rows=1,cols=5);table.style='Table Grid'
for c,t in zip(table.rows[0].cells,['编号','中文标题','English description','页尺寸','核对结果']):run(c.paragraphs[0],t,bold=True)
for i in range(6):
    for c,t in zip(table.add_row().cells,[str(i+1),'文本与矢量','Independent page dimensions','297 × 210 mm','待比较']):run(c.paragraphs[0],t)
doc.add_picture(str(ROOT/'chart.png'),width=Mm(190))
sec=doc.add_section(WD_SECTION_START.NEW_PAGE);section(sec,215.9,279.4)
run(doc.add_paragraph(), '五、Letter 尺寸结论页',size=20,bold=True)
run(doc.add_paragraph(), 'END-OF-REPORT 最后一页。每一个源页面都应对应一个输出页面，成功导出不代表完成打印。')
doc.save(ROOT/'multipage.docx')

# 将字符覆盖和排版复杂度隔离，方便定位后端及 fallback。
cases={
 'chinese': [('纯中文：你好世界，中文搜索复制。简体繁體，标点「测试」。','PingFang SC')],
 'mixed': [('中文 English 123，PDF 搜索 Copy ¥99.50。','PingFang SC')],
 'fonts': [('苹方中文 Arial English','PingFang SC'),('宋体中文 Serif English','Songti SC'),('黑体中文 Mono English','Heiti SC')],
 'missing': [('Fallback 中文缺失字体 English 0123','PlifloMissingCJKXYZ')],
}
for name,lines in cases.items():
    f=Document();section(f.sections[0])
    for text,east in lines:run(f.add_paragraph(),text,'PlifloMissingFontXYZ' if name=='missing' else 'Arial',east,size=16)
    f.save(ROOT/f'{name}.docx')
(ROOT/'manifest.json').write_text(json.dumps({'controlledFixture':True,'description':'Authored business-report DOCX, not a Word-validated golden or user-supplied production document','cases':list(cases),'sentinels':['SEARCH-CJK-001','RECORD-001','RECORD-030','END-OF-REPORT','你好世界','中文搜索复制验证']},ensure_ascii=False,indent=2))
stress=Document();section(stress.sections[0])
for i in range(40):
    if i:stress.add_page_break()
    run(stress.add_paragraph(),f'PAGE-{i+1:03d} 内存与多页压力测试',size=18,bold=True)
    for j in range(8):run(stress.add_paragraph(),f'{j+1}. 中文与 English 0123456789 文档渲染需要保持字体、边距和段落一致。每页独立输出，不代表物理打印完成。')
    if i%10==0:stress.add_picture(str(ROOT/'chart.png'),width=Mm(120))
stress.save(ROOT/'stress-40.docx')
print(ROOT/'multipage.docx')
